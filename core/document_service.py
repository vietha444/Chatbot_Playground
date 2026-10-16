# -*- coding: utf-8 -*-
"""
Document Service - Quản lý Kho Tài Liệu, Tải Lên & Lập Chỉ Mục RAG
Tích hợp toàn diện các thao tác từ:
- 1_Thêm tài liệu.py (PdfReader / PyMuPDF text extraction, RecursiveCharacterTextSplitter, ChromaDB + Ollama bge-m3, appdatabase insert)
- 2_Xóa tài liệu.py (appdatabase delete, ChromaDB collection delete)
- Xem nội dung tài liệu lưu trong PostgreSQL Neon DB (cột doc của bảng appdatabase)
- Phân quyền: Người dùng bình thường chỉ quản lý tài liệu của mình; Admin quản lý toàn bộ.
"""

import os
import sys
import re
import json
import time
import psycopg2

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

from typing import List, Dict, Any, Optional, Tuple
import requests
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename
try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:
    from langchain.text_splitter import RecursiveCharacterTextSplitter
import chromadb
from chromadb.api.types import Documents, EmbeddingFunction, Embeddings
from core.config import Config

# Thư viện đọc tài liệu
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pptx
except ImportError:
    pptx = None


from core.embedding_service import OllamaBgeM3EmbeddingFunction


class DocumentService:
    def __init__(self):
        self.metadata_file = Config.DOCUMENTS_METADATA_FILE
        self.upload_dir = Config.UPLOAD_FOLDER
        os.makedirs(self.upload_dir, exist_ok=True)
        os.makedirs(os.path.dirname(self.metadata_file), exist_ok=True)

        if not os.path.exists(self.metadata_file):
            with open(self.metadata_file, "w", encoding="utf-8") as f:
                json.dump([], f, ensure_ascii=False)

        # Khởi tạo BGE-M3 Embedding Function độc quyền
        self.bge_ef = OllamaBgeM3EmbeddingFunction(
            endpoint=Config.OLLAMA_REMOTE_URL,
            model=Config.EMBEDDING_MODEL
        )

        # Chroma Local Client
        self.local_chroma = chromadb.PersistentClient(path=Config.CHROMA_PERSIST_DIR)

        # Chroma Remote Client (Tailscale / HTTP)
        try:
            self.remote_chroma = chromadb.HttpClient(host=Config.CHROMA_SERVER_URL)
        except Exception:
            self.remote_chroma = None

    def get_db_connection(self):
        return psycopg2.connect(
            dbname=Config.DB_NAME,
            user=Config.DB_USER,
            password=Config.DB_PASSWORD,
            host=Config.DB_HOST,
            port=Config.DB_PORT,
            sslmode="require",
            connect_timeout=6
        )

    # ---------------------- TRÍCH XUẤT VĂN BẢN ----------------------
    def extract_text_from_file(self, file_path: str) -> Tuple[str, int]:
        """
        Trích xuất toàn bộ văn bản từ file (PDF, DOCX, XLSX, PPTX, TXT, Code, Markdown).
        Trả về (full_text, total_pages).
        """
        ext = file_path.split('.')[-1].lower()
        full_text = ""
        total_pages = 1

        if ext == "pdf":
            # Ưu tiên PyMuPDF (fitz) vì nhanh và trích xuất chuẩn
            if fitz is not None:
                try:
                    doc = fitz.open(file_path)
                    total_pages = len(doc)
                    pages_text = []
                    for page in doc:
                        t = page.get_text()
                        if t and t.strip():
                            pages_text.append(t)
                    full_text = "\n\n".join(pages_text)
                    doc.close()
                except Exception as e:
                    print(f"[DocumentService] Lỗi fitz: {e}, chuyển sang PyPDF2...")
                    full_text = ""

            # Fallback sang PyPDF2 (như trong 1_Thêm tài liệu.py)
            if not full_text and PdfReader is not None:
                try:
                    reader = PdfReader(file_path)
                    total_pages = len(reader.pages)
                    pages_text = []
                    for page in reader.pages:
                        extracted = page.extract_text()
                        if extracted and extracted.strip():
                            pages_text.append(extracted)
                    full_text = "\n\n".join(pages_text)
                except Exception as e:
                    print(f"[DocumentService] Lỗi PyPDF2: {e}")

        elif ext in ["docx", "doc"]:
            if docx is not None:
                try:
                    doc = docx.Document(file_path)
                    parts = []
                    for p in doc.paragraphs:
                        if p.text.strip():
                            parts.append(p.text.strip())
                    for tbl in doc.tables:
                        for row in tbl.rows:
                            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                            if any(cells):
                                parts.append(" | ".join(cells))
                    full_text = "\n\n".join(parts)
                    total_pages = max(1, len(full_text) // 2000)
                except Exception as e:
                    print(f"[DocumentService] Lỗi đọc docx: {e}")

        elif ext in ["xlsx", "xls"]:
            if openpyxl is not None:
                try:
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    sheet_texts = []
                    for sheet in wb.worksheets:
                        sheet_lines = [f"### Trang tính: {sheet.title}"]
                        has_data = False
                        for row in sheet.iter_rows(values_only=True):
                            if any(c is not None and str(c).strip() != "" for c in row):
                                row_str = " | ".join([str(c) if c is not None else "" for c in row])
                                sheet_lines.append(row_str)
                                has_data = True
                        if has_data:
                            sheet_texts.append("\n".join(sheet_lines))
                    full_text = "\n\n".join(sheet_texts)
                    total_pages = max(1, len(wb.worksheets))
                except Exception as e:
                    print(f"[DocumentService] Lỗi đọc xlsx: {e}")

        elif ext in ["pptx", "ppt"]:
            if pptx is not None:
                try:
                    prs = pptx.Presentation(file_path)
                    slide_texts = []
                    for idx, slide in enumerate(prs.slides, start=1):
                        texts = []
                        for shape in slide.shapes:
                            if shape.has_text_frame:
                                for p in shape.text_frame.paragraphs:
                                    if p.text.strip():
                                        texts.append(p.text.strip())
                        if texts:
                            slide_texts.append(f"--- Slide {idx} ---\n" + "\n".join(texts))
                    full_text = "\n\n".join(slide_texts)
                    total_pages = max(1, len(prs.slides))
                except Exception as e:
                    print(f"[DocumentService] Lỗi đọc pptx: {e}")

        else:
            # Code và văn bản thuần (txt, md, csv, json, py, js, ts, html, css, c, cpp, java, sql, yaml, ...)
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    full_text = f.read()
                total_pages = max(1, len(full_text) // 2500)
            except Exception as e:
                print(f"[DocumentService] Lỗi đọc file text: {e}")

        return full_text.strip(), total_pages

    # ---------------------- TẢI LÊN & LẬP CHỈ MỤC RAG ĐỒNG BỘ ----------------------
    def upload_and_index_document(
        self,
        file_obj: FileStorage,
        tenvanban: str,
        tencollection: str,
        uploader: str = "user"
    ) -> Dict[str, Any]:
        """
        Thực hiện toàn bộ quy trình đồng bộ chuẩn:
        1. Lưu file vật lý vào thư mục uploads/
        2. Bóc tách văn bản (PDF / DOCX / TXT)
        3. Cắt đoạn văn bản qua RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        4. Ghi toàn bộ nội dung PDF (full_text) vào cột doc của PostgreSQL Neon (bảng appdatabase)
        5. Lập chỉ mục Vector vào ChromaDB với Ollama bge-m3 embeddings
        6. Lưu metadata vào documents.json
        """
        if not file_obj or file_obj.filename == "":
            raise ValueError("File tải lên không hợp lệ.")

        tenvanban = tenvanban.strip() if tenvanban else ""
        if not tenvanban:
            tenvanban = os.path.splitext(file_obj.filename)[0]

        tencollection = tencollection.strip() if tencollection else ""
        if not tencollection:
            clean_name = re.sub(r'[^a-zA-Z0-9_]', '_', tenvanban.lower())[:20].strip('_')
            tencollection = f"col_{clean_name}_{int(time.time()) % 10000}"

        tencollection = re.sub(r'[^a-zA-Z0-9_\-]', '_', tencollection)

        # 1. Lưu file vật lý
        safe_fname = secure_filename(file_obj.filename)
        timestamp = int(time.time())
        saved_filename = f"{timestamp}_{safe_fname}"
        saved_path = os.path.join(self.upload_dir, saved_filename)
        file_obj.save(saved_path)

        # 2. Bóc tách văn bản
        full_text, total_pages = self.extract_text_from_file(saved_path)
        if not full_text:
            raise ValueError("Không thể trích xuất văn bản từ tài liệu này hoặc file rỗng.")

        # 3. Cắt đoạn theo thuật toán chuẩn từ 1_Thêm tài liệu.py
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        chunks = text_splitter.split_text(full_text)
        if not chunks:
            chunks = [full_text]

        # 4. Ghi vào Neon PostgreSQL (Bảng appdatabase, lưu trữ đầy đủ full_text vào cột doc)
        try:
            conn = self.get_db_connection()
            cur = conn.cursor()
            # Xóa bản ghi cũ cùng collection nếu có để tránh trùng lặp
            cur.execute("DELETE FROM appdatabase WHERE tencollection = %s", (tencollection,))
            cur.execute(
                """
                INSERT INTO appdatabase (tenvanban, tencollection, doc, uploader, created_at)
                VALUES (%s, %s, %s, %s, NOW())
                """,
                (tenvanban, tencollection, full_text, uploader)
            )
            conn.commit()
            cur.close()
            conn.close()
            print(f"[DocumentService] Đã lưu vào Neon DB appdatabase: '{tenvanban}' ({len(full_text)} ký tự, uploader: {uploader})")
        except Exception as e:
            print(f"[DocumentService] Lỗi lưu Neon DB appdatabase: {e}")

        # 5. Lập chỉ mục vào ChromaDB với Ollama bge-m3 embeddings (KHÔNG nạp hay dùng all-MiniLM-L6-v2)
        try:
            print(f"[DocumentService] Đang tính bge-m3 embeddings cho {len(chunks)} chunks từ {Config.OLLAMA_REMOTE_URL}...")
            embeddings_list = self.bge_ef(chunks)
            emb_dim = len(embeddings_list[0]) if embeddings_list else 0
            print(f"[DocumentService] Đã tính xong {len(embeddings_list)} embeddings (dim: {emb_dim}) với mô hình bge-m3")

            # 5a. Local ChromaDB với bge_ef
            local_col = self.local_chroma.get_or_create_collection(
                name=tencollection,
                embedding_function=self.bge_ef
            )
            doc_ids = [f"{tencollection}_chunk_{i}_{timestamp}" for i in range(len(chunks))]
            metadatas = [{"source": tenvanban, "collection": tencollection, "uploader": uploader, "page": 1} for _ in chunks]

            batch_size = 50
            for i in range(0, len(chunks), batch_size):
                end_i = min(i + batch_size, len(chunks))
                local_col.add(
                    documents=chunks[i:end_i],
                    embeddings=embeddings_list[i:end_i],
                    metadatas=metadatas[i:end_i],
                    ids=doc_ids[i:end_i]
                )
            print(f"[DocumentService] Đã nạp thành công {len(chunks)} chunks vào Local Chroma '{tencollection}' (1024-dim bge-m3)")

            # 5b. Remote Chroma nếu có
            if self.remote_chroma:
                try:
                    remote_col = self.remote_chroma.get_or_create_collection(
                        name=tencollection,
                        embedding_function=self.bge_ef
                    )
                    for i in range(0, len(chunks), batch_size):
                        end_i = min(i + batch_size, len(chunks))
                        remote_col.add(
                            documents=chunks[i:end_i],
                            embeddings=embeddings_list[i:end_i],
                            metadatas=metadatas[i:end_i],
                            ids=doc_ids[i:end_i]
                        )
                    print(f"[DocumentService] Đã đồng bộ {len(chunks)} chunks sang Remote Chroma '{tencollection}'")
                except Exception as re_err:
                    print(f"[DocumentService] Remote Chroma không khả dụng lúc này: {re_err}")
        except Exception as c_err:
            print(f"[DocumentService] Lỗi nạp ChromaDB: {c_err}")

        # 6. Ghi thông tin file vào metadata nội bộ
        doc_id = f"doc_{timestamp}_{os.urandom(3).hex()}"
        doc_record = {
            "id": doc_id,
            "filename": tenvanban,
            "saved_filename": saved_filename,
            "file_path": saved_path,
            "relative_url": f"/api/documents/{doc_id}/file",
            "file_size": os.path.getsize(saved_path),
            "total_pages": total_pages,
            "collection": tencollection,
            "chunks_count": len(chunks),
            "uploader": uploader,
            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "has_physical_file": True,
            "text_preview": full_text[:300] + ("..." if len(full_text) > 300 else "")
        }

        metadata_list = self._read_metadata()
        metadata_list.insert(0, doc_record)
        self._write_metadata(metadata_list)

        return doc_record

    # ---------------------- LIỆT KÊ TÀI LIỆU CÓ PHÂN QUYỀN ----------------------
    def list_documents(self, username: Optional[str] = None, role: str = "user") -> List[Dict[str, Any]]:
        """
        Lấy danh sách tài liệu từ Neon PostgreSQL appdatabase và đồng bộ file xem trước.
        - Nếu role == 'admin': Xem toàn bộ tài liệu của tất cả người dùng.
        - Nếu role == 'user': Chỉ xem tài liệu của chính mình (uploader == username) hoặc tài liệu mẫu (system).
        """
        is_admin = (role == "admin" or (username and username.lower() == "admin"))
        local_meta = {d.get("collection"): d for d in self._read_metadata() if d.get("collection")}
        docs = []

        try:
            conn = self.get_db_connection()
            cur = conn.cursor()
            if is_admin:
                cur.execute(
                    """
                    SELECT tenvanban, tencollection, length(doc), uploader, created_at, substring(doc from 1 for 250)
                    FROM appdatabase
                    ORDER BY created_at DESC NULLS LAST
                    """
                )
            else:
                user_filter = username if username else "system"
                cur.execute(
                    """
                    SELECT tenvanban, tencollection, length(doc), uploader, created_at, substring(doc from 1 for 250)
                    FROM appdatabase
                    WHERE LOWER(uploader) = LOWER(%s) OR uploader = 'system' OR uploader IS NULL
                    ORDER BY created_at DESC NULLS LAST
                    """,
                    (user_filter,)
                )
            
            rows = cur.fetchall()
            cur.close()
            conn.close()

            for r in rows:
                tenvanban = str(r[0]) if r[0] else "Không tên"
                tencollection = str(r[1]) if r[1] else ""
                doc_len = r[2] if r[2] else 0
                uploader = str(r[3]) if r[3] else "system"
                created_at = str(r[4])[:19] if r[4] else "2026-10-15 00:00:00"
                text_preview = str(r[5]) if r[5] else ""

                # Ghép với metadata file xem trước (nếu có)
                meta_item = local_meta.get(tencollection, {})
                doc_id = meta_item.get("id", f"neon_{tencollection}")
                has_physical_file = bool(meta_item.get("file_path") and os.path.exists(meta_item.get("file_path")))
                file_url = f"/api/documents/{doc_id}/file"

                is_owner = (username and username.lower() == uploader.lower()) or (uploader == "system" and is_admin)
                can_delete = is_admin or is_owner

                docs.append({
                    "id": doc_id,
                    "filename": tenvanban,
                    "collection": tencollection,
                    "file_size": doc_len,
                    "total_pages": meta_item.get("total_pages", max(1, doc_len // 2000)),
                    "chunks_count": meta_item.get("chunks_count", max(1, doc_len // 800)),
                    "uploader": uploader,
                    "created_at": created_at,
                    "is_owner": is_owner,
                    "can_delete": can_delete,
                    "has_physical_file": has_physical_file,
                    "relative_url": file_url,
                    "text_preview": text_preview
                })

        except Exception as e:
            print(f"[DocumentService] Lỗi truy vấn Neon DB appdatabase: {e}")
            all_meta = self._read_metadata()
            for m in all_meta:
                m_uploader = m.get("uploader", "system")
                if is_admin or (username and m_uploader.lower() == username.lower()) or m_uploader == "system":
                    m_copy = dict(m)
                    m_copy["is_owner"] = (username and username.lower() == m_uploader.lower())
                    m_copy["can_delete"] = is_admin or m_copy["is_owner"]
                    m_copy["has_physical_file"] = bool(m.get("file_path") and os.path.exists(m.get("file_path")))
                    docs.append(m_copy)

        return docs

    # ---------------------- LẤY NỘI DUNG VĂN BẢN TỪ NEON DB ĐỂ XEM ----------------------
    def get_document_content(
        self,
        doc_id_or_col: Optional[str] = None,
        collection: Optional[str] = None,
        filename: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Lấy nội dung toàn văn và metadata của tài liệu:
        - Tìm kiếm trong PostgreSQL Neon DB (cột doc của bảng appdatabase) bằng collection, filename hoặc id.
        - Đọc từ file vật lý trên đĩa (nếu là file mới tải lên).
        - Đảm bảo LUÔN trả về dictionary hợp lệ (char_count, content...), không bao giờ trả về None để tránh lỗi 404/HTML.
        """
        target_id = str(doc_id_or_col or "").strip()
        target_col = str(collection or "").strip()
        target_fn = str(filename or "").strip()

        clean_id = target_id.replace("neon_", "").strip() if target_id else ""
        search_terms = list(dict.fromkeys([t for t in [clean_id, target_col, target_fn, target_id] if t]))

        # 1. Tra cứu trong metadata local để kiểm tra file vật lý
        meta = None
        for d in self._read_metadata():
            d_id = d.get("id", "")
            d_col = d.get("collection", "")
            d_fn = d.get("filename", "")
            if any(term in (d_id, d_col, d_fn) for term in search_terms):
                meta = d
                break

        # 2. Truy vấn trực tiếp PostgreSQL Neon DB appdatabase
        neon_doc = None
        try:
            conn = self.get_db_connection()
            cur = conn.cursor()
            for term in search_terms:
                cur.execute(
                    """
                    SELECT tenvanban, tencollection, doc, uploader, created_at
                    FROM appdatabase
                    WHERE tencollection = %s OR tenvanban = %s OR LOWER(tencollection) = LOWER(%s) OR LOWER(tenvanban) = LOWER(%s)
                    LIMIT 1
                    """,
                    (term, term, term, term)
                )
                row = cur.fetchone()
                if row:
                    neon_doc = {
                        "filename": str(row[0]) if row[0] else (target_fn or term),
                        "collection": str(row[1]) if row[1] else (target_col or term),
                        "content": str(row[2]) if row[2] else "",
                        "uploader": str(row[3]) if row[3] else "system",
                        "created_at": str(row[4])[:19] if row[4] else ""
                    }
                    break
            cur.close()
            conn.close()
        except Exception as e:
            print(f"[DocumentService] Lỗi truy vấn Neon DB doc content: {e}")

        # Trường hợp 1: Có dữ liệu text trong Neon DB (Hầu hết các tài liệu)
        if neon_doc and neon_doc.get("content"):
            has_file = False
            file_path = ""
            if meta and meta.get("file_path") and os.path.exists(meta["file_path"]):
                has_file = True
                file_path = meta["file_path"]

            doc_repr_id = target_id or f"neon_{neon_doc['collection']}"
            return {
                "id": doc_repr_id,
                "filename": neon_doc["filename"],
                "collection": neon_doc["collection"],
                "uploader": neon_doc["uploader"],
                "created_at": neon_doc["created_at"],
                "char_count": len(neon_doc["content"]),
                "word_count": len(neon_doc["content"].split()),
                "content": neon_doc["content"],
                "has_physical_file": has_file,
                "file_url": f"/api/documents/{doc_repr_id}/file" if has_file else None
            }

        # Trường hợp 2: Có file vật lý trên máy
        if meta and meta.get("file_path") and os.path.exists(meta["file_path"]):
            full_text, _ = self.extract_text_from_file(meta["file_path"])
            doc_repr_id = meta.get("id", target_id or f"doc_{int(time.time())}")
            return {
                "id": doc_repr_id,
                "filename": meta.get("filename", target_fn or "Tài liệu"),
                "collection": meta.get("collection", target_col or "Mặc định"),
                "uploader": meta.get("uploader", "system"),
                "created_at": meta.get("created_at", ""),
                "char_count": len(full_text),
                "word_count": len(full_text.split()),
                "content": full_text,
                "has_physical_file": True,
                "file_url": f"/api/documents/{doc_repr_id}/file"
            }

        # Trường hợp 3: Neon DB có hàng nhưng cột doc rỗng
        if neon_doc:
            doc_repr_id = target_id or f"neon_{neon_doc['collection']}"
            return {
                "id": doc_repr_id,
                "filename": neon_doc["filename"],
                "collection": neon_doc["collection"],
                "uploader": neon_doc["uploader"],
                "created_at": neon_doc["created_at"],
                "char_count": 0,
                "word_count": 0,
                "content": "(Tài liệu này chưa có nội dung văn bản trong CSDL Neon DB.)",
                "has_physical_file": False,
                "file_url": None
            }

        # Trường hợp 4: Fallback không tìm thấy - vẫn trả về dict hợp lệ thay vì None để tránh 404
        fallback_col = target_col or clean_id or "Mặc định"
        fallback_fn = target_fn or target_col or clean_id or "Tài liệu"
        return {
            "id": target_id or f"neon_{fallback_col}",
            "filename": fallback_fn,
            "collection": fallback_col,
            "uploader": "system",
            "created_at": "",
            "char_count": 0,
            "word_count": 0,
            "content": f"(Chưa tìm thấy nội dung văn bản chi tiết cho tài liệu '{fallback_fn}'. Bạn có thể tải lại file này vào hệ thống để cập nhật nội dung.)",
            "has_physical_file": False,
            "file_url": None
        }

    # ---------------------- XÓA TÀI LIỆU CÓ PHÂN QUYỀN ----------------------
    def delete_document_by_name(
        self,
        tenvanban: str,
        tencollection: str,
        requested_by: str,
        role: str = "user"
    ) -> Tuple[bool, str]:
        """
        Xóa tài liệu khỏi appdatabase và ChromaDB (theo chuẩn 2_Xóa tài liệu.py).
        Kiểm tra quyền:
        - Admin: Xóa được mọi tài liệu.
        - User: Chỉ xóa được tài liệu mà uploader == requested_by.
        """
        is_admin = (role == "admin" or requested_by.lower() == "admin")

        try:
            conn = self.get_db_connection()
            cur = conn.cursor()

            # Kiểm tra uploader của tài liệu
            cur.execute(
                "SELECT uploader FROM appdatabase WHERE tenvanban = %s AND tencollection = %s",
                (tenvanban, tencollection)
            )
            row = cur.fetchone()

            if not row:
                cur.close()
                conn.close()
                return False, f"Không tìm thấy tài liệu '{tenvanban}' trong CSDL."

            doc_uploader = row[0] or "system"

            if not is_admin and doc_uploader.lower() != requested_by.lower():
                cur.close()
                conn.close()
                return False, "Bạn không có quyền xóa tài liệu của người khác."

            # 1. Xóa khỏi PostgreSQL Neon appdatabase (theo 2_Xóa tài liệu.py)
            cur.execute(
                "DELETE FROM appdatabase WHERE tenvanban = %s AND tencollection = %s",
                (tenvanban, tencollection)
            )
            conn.commit()
            cur.close()
            conn.close()
            print(f"[DocumentService] Đã xóa '{tenvanban}' ({tencollection}) khỏi Neon DB")

            # 2. Xóa collection trong ChromaDB
            try:
                self.local_chroma.delete_collection(name=tencollection)
                print(f"[DocumentService] Đã xóa Local Chroma collection '{tencollection}'")
            except Exception:
                pass

            if self.remote_chroma:
                try:
                    self.remote_chroma.delete_collection(name=tencollection)
                    print(f"[DocumentService] Đã xóa Remote Chroma collection '{tencollection}'")
                except Exception:
                    pass

            # 3. Cập nhật metadata file local và xóa file vật lý nếu có
            all_meta = self._read_metadata()
            new_meta = []
            for m in all_meta:
                if m.get("collection") == tencollection or m.get("filename") == tenvanban:
                    fpath = m.get("file_path")
                    if fpath and os.path.exists(fpath):
                        try:
                            os.remove(fpath)
                        except Exception:
                            pass
                else:
                    new_meta.append(m)
            self._write_metadata(new_meta)

            return True, f"Đã xóa thành công tài liệu '{tenvanban}' khỏi hệ thống!"
        except Exception as e:
            return False, f"Lỗi trong quá trình xóa: {str(e)}"

    def get_document(self, doc_id: str) -> Optional[Dict[str, Any]]:
        for d in self._read_metadata():
            if d.get("id") == doc_id or d.get("collection") == doc_id.replace("neon_", ""):
                return d
        return None

    def _read_metadata(self) -> List[Dict[str, Any]]:
        try:
            with open(self.metadata_file, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []

    def _write_metadata(self, data: List[Dict[str, Any]]) -> None:
        try:
            with open(self.metadata_file, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass
