"""
Bộ bóc tách và trích xuất nội dung từ các định dạng tài liệu Office & PDF
Hỗ trợ: .docx, .xlsx, .pptx, .pdf, .csv, .txt, .md
"""
import os
import json
import logging
from typing import Dict, Any, List, Optional

logger = logging.getLogger(__name__)

def extract_docx_content(file_path: str) -> Dict[str, Any]:
    """Trích xuất đoạn văn, tiêu đề và bảng biểu từ file Word (.docx)"""
    import docx
    
    doc = docx.Document(file_path)
    paragraphs = []
    headings = []
    tables_data = []
    full_text_parts = []
    
    # 1. Trích xuất đoạn văn và tiêu đề
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        
        is_heading = p.style.name.startswith("Heading") or p.style.name in ["Title", "Subtitle"]
        p_info = {
            "text": text,
            "style": p.style.name,
            "is_heading": is_heading
        }
        paragraphs.append(p_info)
        full_text_parts.append(text)
        if is_heading:
            headings.append(text)
            
    # 2. Trích xuất bảng biểu
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            table_rows.append(row_cells)
        if table_rows:
            tables_data.append({
                "table_index": t_idx + 1,
                "rows": table_rows
            })
            # Biểu diễn bảng dạng Markdown cho LLM dễ đọc
            if len(table_rows) > 0:
                header = table_rows[0]
                md_table = "\n| " + " | ".join(header) + " |\n"
                md_table += "| " + " | ".join(["---"] * len(header)) + " |\n"
                for r in table_rows[1:]:
                    md_table += "| " + " | ".join(r) + " |\n"
                full_text_parts.append(f"\n[BẢNG {t_idx + 1}]:\n{md_table}")

    return {
        "file_type": "docx",
        "file_name": os.path.basename(file_path),
        "total_paragraphs": len(paragraphs),
        "total_tables": len(tables_data),
        "headings": headings,
        "paragraphs": paragraphs,
        "tables": tables_data,
        "full_text": "\n\n".join(full_text_parts)
    }


def extract_xlsx_content(file_path: str) -> Dict[str, Any]:
    """Trích xuất các sheet, dòng, cột và công thức từ file Excel (.xlsx)"""
    import openpyxl
    
    # Mở workbook ở chế độ data_only=False để đọc công thức thực tế (=SUM...)
    wb_formulas = openpyxl.load_workbook(file_path, data_only=False)
    # Mở ở chế độ data_only=True để đọc giá trị đã tính toán (nếu có cache)
    try:
        wb_values = openpyxl.load_workbook(file_path, data_only=True)
    except Exception:
        wb_values = None

    sheets_data = []
    full_text_parts = []
    
    for sheet_name in wb_formulas.sheetnames:
        ws_f = wb_formulas[sheet_name]
        ws_v = wb_values[sheet_name] if wb_values and sheet_name in wb_values.sheetnames else None
        
        rows_data = []
        formulas_found = []
        max_row = min(ws_f.max_row, 300)  # Giới hạn 300 dòng để tránh tràn context
        max_col = min(ws_f.max_column, 40)
        
        sheet_text = [f"=== SHEET: {sheet_name} ==="]
        
        for r_idx in range(1, max_row + 1):
            row_vals = []
            for c_idx in range(1, max_col + 1):
                cell_f = ws_f.cell(row=r_idx, column=c_idx)
                val_f = cell_f.value
                
                # Kiểm tra xem ô có chứa công thức không
                val_str = str(val_f) if val_f is not None else ""
                if isinstance(val_f, str) and val_f.startswith("="):
                    cell_coord = cell_f.coordinate
                    formulas_found.append({
                        "cell": cell_coord,
                        "formula": val_f
                    })
                    val_str = f"{val_f}"
                row_vals.append(val_str)
            
            # Chỉ lưu dòng nếu có ít nhất 1 ô có dữ liệu
            if any(cell_val.strip() for cell_val in row_vals):
                rows_data.append(row_vals)
                sheet_text.append(" | ".join(row_vals))
                
        sheets_data.append({
            "sheet_name": sheet_name,
            "total_rows": len(rows_data),
            "rows": rows_data,
            "formulas": formulas_found
        })
        full_text_parts.append("\n".join(sheet_text))

    return {
        "file_type": "xlsx",
        "file_name": os.path.basename(file_path),
        "sheets": sheets_data,
        "total_sheets": len(sheets_data),
        "full_text": "\n\n".join(full_text_parts)
    }


def extract_pptx_content(file_path: str) -> Dict[str, Any]:
    """Trích xuất các Slide, tiêu đề, nội dung và ghi chú từ file PowerPoint (.pptx)"""
    import pptx
    
    prs = pptx.Presentation(file_path)
    slides_data = []
    full_text_parts = []
    
    for s_idx, slide in enumerate(prs.slides):
        slide_title = ""
        slide_texts = []
        tables_in_slide = []
        
        # Đọc tiêu đề nếu có
        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()
            
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for para in shape.text_frame.paragraphs:
                    p_text = para.text.strip()
                    if p_text:
                        slide_texts.append(p_text)
            elif shape.has_table:
                table = shape.table
                t_rows = []
                for row in table.rows:
                    r_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    t_rows.append(r_cells)
                tables_in_slide.append(t_rows)

        # Ghi chú diễn giả (notes)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slide_info = {
            "slide_number": s_idx + 1,
            "title": slide_title,
            "content": slide_texts,
            "tables": tables_in_slide,
            "notes": notes
        }
        slides_data.append(slide_info)
        
        # Biểu diễn text cho slide
        s_text = f"--- SLIDE {s_idx + 1}: {slide_title} ---\n"
        if slide_texts:
            s_text += "\n".join([f"- {t}" for t in slide_texts]) + "\n"
        if tables_in_slide:
            for t in tables_in_slide:
                s_text += "\n" + "\n".join([" | ".join(row) for row in t]) + "\n"
        if notes:
            s_text += f"[Ghi chú]: {notes}\n"
        full_text_parts.append(s_text)

    return {
        "file_type": "pptx",
        "file_name": os.path.basename(file_path),
        "total_slides": len(slides_data),
        "slides": slides_data,
        "full_text": "\n\n".join(full_text_parts)
    }


def extract_pdf_content(file_path: str) -> Dict[str, Any]:
    """Trích xuất nội dung văn bản và cấu trúc trang từ file PDF bằng PyMuPDF (fitz)"""
    import fitz  # PyMuPDF
    
    doc = fitz.open(file_path)
    pages_data = []
    full_text_parts = []
    has_scanned_page = False
    
    for p_idx, page in enumerate(doc):
        p_num = p_idx + 1
        text = page.get_text("text").strip()
        
        # Nếu trang hầu như không có text nhưng có ảnh thì có thể là bản scan
        if len(text) < 30 and len(page.get_images()) > 0:
            has_scanned_page = True
            
        pages_data.append({
            "page_number": p_num,
            "text": text,
            "char_count": len(text)
        })
        if text:
            full_text_parts.append(f"--- TRANG {p_num} ---\n{text}")
            
    doc.close()
    
    return {
        "file_type": "pdf",
        "file_name": os.path.basename(file_path),
        "total_pages": len(pages_data),
        "has_scanned_page": has_scanned_page,
        "pages": pages_data,
        "full_text": "\n\n".join(full_text_parts)
    }


def extract_plain_text(file_path: str) -> Dict[str, Any]:
    """Trích xuất file text thuần (.txt, .md, .csv, .json)"""
    encodings = ["utf-8", "utf-8-sig", "cp1258", "latin1"]
    content = ""
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue
            
    return {
        "file_type": "text",
        "file_name": os.path.basename(file_path),
        "full_text": content
    }


def extract_any_document(file_path: str) -> Dict[str, Any]:
    """Hàm bóc tách tổng quát tự động nhận diện phần mở rộng của file"""
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")
        
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".docx":
        return extract_docx_content(file_path)
    elif ext in [".xlsx", ".xls"]:
        return extract_xlsx_content(file_path)
    elif ext in [".pptx", ".ppt"]:
        return extract_pptx_content(file_path)
    elif ext == ".pdf":
        return extract_pdf_content(file_path)
    elif ext in [".txt", ".md", ".csv", ".json", ".log"]:
        return extract_plain_text(file_path)
    else:
        # Thử đọc dạng text thuần nếu là định dạng khác
        try:
            return extract_plain_text(file_path)
        except Exception as e:
            logger.warning(f"Không hỗ trợ bóc tách định dạng {ext}: {e}")
            return {
                "file_type": "unknown",
                "file_name": os.path.basename(file_path),
                "full_text": f"[Định dạng {ext} chưa được hỗ trợ bóc tách chi tiết]"
            }
