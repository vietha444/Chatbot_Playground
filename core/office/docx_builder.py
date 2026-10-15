"""
Module tạo và định dạng tài liệu Microsoft Word (.docx) chuyên nghiệp
Hỗ trợ cả định dạng cấu trúc JSON và chuyển đổi Markdown sang Word chuẩn thể thức.
"""
import os
import re
from typing import Dict, Any, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def _set_cell_background(cell, hex_color: str):
    """Đặt màu nền cho ô trong bảng (ví dụ: '1E3A8A' hoặc 'F1F5F9')"""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề padding bên trong ô của bảng"""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def _set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Đặt đường viền tinh tế cho toàn bộ bảng"""
    tbl_pr = table._element.xpath('w:tblPr')
    if tbl_pr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tbl_pr[0].append(borders)

def init_styled_document() -> docx.Document:
    """Khởi tạo Document Word với căn lề và style mặc định chuẩn Việt Nam"""
    doc = docx.Document()
    
    # Căn lề chuẩn văn bản: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # ~2.0 cm
        section.bottom_margin = Inches(0.79)  # ~2.0 cm
        section.left_margin = Inches(1.18)    # ~3.0 cm
        section.right_margin = Inches(0.79)   # ~2.0 cm
        section.page_width = Inches(8.27)     # A4
        section.page_height = Inches(11.69)
        
    # Thiết lập style Normal mặc định
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37) # Màu xám đậm lịch sự
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)
    
    return doc

def build_docx_from_data(data: Dict[str, Any], output_path: str) -> str:
    """
    Tạo file Word từ cấu trúc dữ liệu JSON chi tiết.
    data format:
    {
        "title": "...",
        "subtitle": "...",
        "organization": "...",
        "sections": [
            {"heading": "...", "level": 1, "content": ["..."], "type": "paragraphs"},
            {"heading": "...", "level": 2, "items": ["..."], "type": "bullets"},
            {"heading": "...", "type": "table", "headers": [...], "rows": [[...]]},
            {"type": "callout", "text": "..."}
        ],
        "signatures": [{"title": "Người lập", "name": "..."}, {"title": "Thủ trưởng", "name": "..."}]
    }
    """
    doc = init_styled_document()
    
    # Header cơ quan / Tiêu ngữ (nếu có)
    org_name = data.get("organization") or data.get("agency")
    if org_name:
        p_org = doc.add_paragraph()
        p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_org = p_org.add_run(org_name.upper() + "\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n---------------")
        r_org.font.name = "Times New Roman"
        r_org.font.size = Pt(11)
        r_org.bold = True
        p_org.paragraph_format.space_after = Pt(14)

    # Tiêu đề tài liệu (Title)
    title = data.get("title", "TÀI LIỆU VĂN BẢN")
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title.upper())
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy Blue
    p_title.paragraph_format.space_after = Pt(4)
    
    # Tiêu đề phụ (Subtitle)
    if data.get("subtitle"):
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(data["subtitle"])
        r_sub.font.name = "Times New Roman"
        r_sub.font.size = Pt(12)
        r_sub.italic = True
        r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        p_sub.paragraph_format.space_after = Pt(16)
    else:
        p_title.paragraph_format.space_after = Pt(16)

    # Các phần nội dung (Sections)
    sections = data.get("sections", [])
    for sec in sections:
        heading_text = sec.get("heading")
        level = sec.get("level", 1)
        
        if heading_text:
            p_head = doc.add_paragraph()
            r_head = p_head.add_run(heading_text)
            r_head.font.name = "Times New Roman"
            r_head.bold = True
            
            if level == 1:
                r_head.font.size = Pt(14)
                r_head.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
                p_head.paragraph_format.space_before = Pt(12)
                p_head.paragraph_format.space_after = Pt(6)
            else:
                r_head.font.size = Pt(13)
                r_head.italic = (level > 2)
                r_head.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
                p_head.paragraph_format.space_before = Pt(8)
                p_head.paragraph_format.space_after = Pt(4)
                
        sec_type = sec.get("type", "paragraphs")
        
        # 1. Đoạn văn
        if sec_type == "paragraphs":
            contents = sec.get("content", [])
            if isinstance(contents, str):
                contents = [contents]
            for para_text in contents:
                p_body = doc.add_paragraph()
                p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_body.paragraph_format.first_line_indent = Inches(0.4)
                p_body.paragraph_format.space_after = Pt(6)
                r = p_body.add_run(para_text)
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                
        # 2. Danh sách gạch đầu dòng (Bullets)
        elif sec_type == "bullets":
            items = sec.get("items", [])
            for item in items:
                p_bullet = doc.add_paragraph(style='List Bullet')
                p_bullet.paragraph_format.space_after = Pt(3)
                r = p_bullet.add_run(item)
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                
        # 3. Bảng biểu (Table)
        elif sec_type == "table":
            headers = sec.get("headers", [])
            rows = sec.get("rows", [])
            if headers or rows:
                col_count = len(headers) if headers else (len(rows[0]) if rows else 1)
                table = doc.add_table(rows=0, cols=col_count)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                _set_table_borders(table)
                
                # Header row
                if headers:
                    hdr_row = table.add_row()
                    for c_idx, h_text in enumerate(headers):
                        cell = hdr_row.cells[c_idx]
                        _set_cell_background(cell, "1E3A8A") # Dark Navy Header
                        _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_after = Pt(0)
                        r = p.add_run(str(h_text))
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(12)
                        r.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # Trắng
                
                # Data rows
                for r_idx, r_data in enumerate(rows):
                    row = table.add_row()
                    bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF" # Zebra
                    for c_idx, val in enumerate(r_data):
                        if c_idx < col_count:
                            cell = row.cells[c_idx]
                            _set_cell_background(cell, bg_color)
                            _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_after = Pt(0)
                            # Căn giữa nếu là số thứ tự hoặc số ngắn
                            val_str = str(val).strip()
                            if len(val_str) <= 4 or val_str.isdigit():
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            r = p.add_run(val_str)
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(12)
                            
                # Khoảng trống sau bảng
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(10)
                
        # 4. Hộp ghi chú / Khung thông tin nổi bật (Callout)
        elif sec_type == "callout":
            callout_text = sec.get("text", "")
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            _set_cell_background(cell, "F1F5F9")
            _set_cell_margins(cell, top=140, bottom=140, left=200, right=150)
            # Viền xanh bên trái
            tc_pr = cell._element.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
                f'<w:top w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'</w:tcBorders>'
            )
            tc_pr.append(borders)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(callout_text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.italic = True
            
            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_after = Pt(8)

    # Chữ ký (Signatures)
    signatures = data.get("signatures", [])
    if signatures:
        p_sig_space = doc.add_paragraph()
        p_sig_space.paragraph_format.space_before = Pt(20)
        
        sig_table = doc.add_table(rows=1, cols=len(signatures))
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, sig in enumerate(signatures):
            cell = sig_table.cell(0, idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(60) # Chừa khoảng trống ký tên
            r_title = p.add_run(sig.get("title", "NGƯỜI KÝ").upper() + "\n")
            r_title.bold = True
            r_title.font.name = "Times New Roman"
            r_sub = p.add_run("(Ký, ghi rõ họ tên)\n\n\n\n")
            r_sub.italic = True
            r_sub.font.size = Pt(10)
            r_sub.font.name = "Times New Roman"
            if sig.get("name"):
                r_name = p.add_run(sig.get("name"))
                r_name.bold = True
                r_name.font.name = "Times New Roman"

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path


from .utils import robust_extract_json, clean_raw_llm_text

def build_docx_from_markdown(markdown_text: str, output_path: str, title: str = "") -> str:
    """
    Chuyển đổi văn bản Markdown từ đầu ra LLM thành tài liệu Word (.docx) chuẩn đẹp.
    Tự động nhận diện tiêu đề (#, ##, ###), bảng markdown (| col | col |), và gạch đầu dòng (- / *).
    Nếu văn bản truyền vào thực chất là cấu trúc JSON, tự động phân tích và dùng build_docx_from_data.
    """
    # Guard 1: Kiểm tra xem có phải JSON không
    parsed_json = robust_extract_json(markdown_text)
    if parsed_json and isinstance(parsed_json, dict) and ("sections" in parsed_json or "title" in parsed_json):
        return build_docx_from_data(parsed_json, output_path)

    # Guard 2: Dọn dẹp code fences hoặc ký tự JSON rác nếu có
    markdown_text = clean_raw_llm_text(markdown_text)

    doc = init_styled_document()
    
    # Thêm tiêu đề chính nếu có
    if title:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_title = p_title.add_run(title.upper())
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        p_title.paragraph_format.space_after = Pt(16)
        
    lines = markdown_text.strip().split("\n")
    i = 0
    in_table = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        
        # Kiểm tra bảng markdown (|...|)
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # Kết thúc bảng, render bảng
            _render_markdown_table(doc, table_lines)
            table_lines = []
            in_table = False
            
        if not stripped:
            i += 1
            continue
            
        # Tiêu đề Level 1
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(15)
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        # Tiêu đề Level 2
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[3:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(13.5)
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        # Tiêu đề Level 3
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[4:].strip())
            r.bold = True
            r.italic = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
        # Danh sách gạch đầu dòng
        elif stripped.startswith("- ") or stripped.startswith("* ") or stripped.startswith("+ "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            _add_formatted_runs(p, stripped[2:].strip())
        # Danh sách số (1. 2. ...)
        elif re.match(r'^\d+\.\s+', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            _add_formatted_runs(p, stripped)
        # Khung trích dẫn (> ...)
        elif stripped.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            _set_cell_background(cell, "F1F5F9")
            _set_cell_margins(cell, top=100, bottom=100, left=150, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(stripped[2:].strip())
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.italic = True
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        # Đoạn văn thường
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
            _add_formatted_runs(p, stripped)
            
        i += 1
        
    # Render bảng còn sót ở cuối file nếu có
    if in_table and table_lines:
        _render_markdown_table(doc, table_lines)
        
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path

def _add_formatted_runs(paragraph, text: str):
    """Phân tách các thẻ in đậm **text** hoặc in nghiêng *text* trong đoạn văn"""
    # Tách token in đậm **...**
    pattern = r'(\*\*.*?\*\*|\*.*?\*)'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        else:
            r = paragraph.add_run(token)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

def _render_markdown_table(doc, lines: List[str]):
    """Vẽ bảng Word từ các dòng bảng Markdown"""
    rows_data = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        # Bỏ qua dòng kẻ ngang (|---|---|)
        if all(re.match(r'^-+$', c) for c in cells):
            continue
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    col_count = max(len(r) for r in rows_data)
    table = doc.add_table(rows=0, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    
    # Dòng Header
    hdr_row = table.add_row()
    for c_idx, text in enumerate(rows_data[0]):
        cell = hdr_row.cells[c_idx]
        _set_cell_background(cell, "1E3A8A")
        _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Các dòng dữ liệu
    for r_idx, r_cells in enumerate(rows_data[1:]):
        row = table.add_row()
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, text in enumerate(r_cells):
            if c_idx < col_count:
                cell = row.cells[c_idx]
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                val_str = text.strip()
                if len(val_str) <= 4 or val_str.isdigit():
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val_str)
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
